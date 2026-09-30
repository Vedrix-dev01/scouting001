import json, glob, re, sys, os
from collections import Counter
from urllib.parse import urlparse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
D=os.path.dirname(os.path.abspath(__file__))
OUT=sys.argv[sys.argv.index("--out")+1] if "--out" in sys.argv else os.path.join(D,"Germany_US_UK_Invoice_AI_Automation_Prospects_1000.xlsx")
COLS=["Prospect Name","Company","First Name","Job Title","Email","Country","State","City","Company Type","Industry",
"Invoice Product or Service","Accounting Platform","ERP","Payment Platform","Company Size","Website","LinkedIn","Buying Intent",
"Automation Opportunity","Personalization Detail","Personalization Reason","Personalized Opening","Subject Line","Email Body",
"Source URL","Evidence","Evidence Type","Research Confidence","Lead Status","Send Status","Date Sent","Notes"]
F=["prospect_name","company","first_name","job_title","email","country","state","city","company_type","industry","invoice_product",
"accounting_platform","erp","payment_platform","company_size","website","linkedin","buying_intent","automation_opportunity",
"personalization_detail","personalization_reason","personalized_opening","subject_line","email_body","source_url","evidence",
"evidence_type","research_confidence"]
OK_COUNTRIES={"United States","United Kingdom","Germany","Netherlands","France","Canada","Australia","Switzerland","Austria","Belgium",
"Ireland","Sweden","Denmark","Norway","Finland","Spain","Italy","Portugal","Poland","Luxembourg","Estonia","Czech Republic","Czechia",
"Lithuania","Latvia","Iceland","Greece","Romania","Hungary","Slovenia","Slovakia","Croatia","Bulgaria","Malta","Cyprus","Serbia","Ukraine"}
CN={"USA":"United States","US":"United States","UK":"United Kingdom","Great Britain":"United Kingdom","England":"United Kingdom",
"Scotland":"United Kingdom","Deutschland":"Germany","The Netherlands":"Netherlands","Holland":"Netherlands","Schweiz":"Switzerland","Österreich":"Austria"}
BANNED=["seamless","tailored","stunning","revolutioni","game changer","game-changer","cutting edge","cutting-edge","unlock","supercharge","transform your business","leverage"]
PERSONAL={"gmail.com","googlemail.com","hotmail.com","outlook.com","yahoo.com","icloud.com","me.com","protonmail.com","proton.me","live.com","gmx.de","web.de","aol.com","t-online.de","yahoo.co.uk","mac.com"}
EXCL={}
if os.path.exists(os.path.join(D,"curation_inv.py")):
    sys.path.insert(0,D); from curation_inv import EXCLUDE as EXCL
EXCL={k.lower():v for k,v in EXCL.items()}
REVIEW={}
try:
    from curation_inv import REVIEW
except Exception: pass
TECH={"github","npm","devops","infra","eng-leadership","system","oss","opensource","open-source","developers","dev","tech","it","webmaster","noreply","npm-publish","eo-sdk","engineering","desarrollo","tpaccount"}
EM=re.compile(r"^[A-Za-z0-9._%+\-]+@[A-Za-z0-9.\-]+\.[A-Za-z]{2,}$")
def dom(u):
    u=(u or "").strip()
    if not u: return ""
    return urlparse(u if "://" in u else "http://"+u).netloc.lower().removeprefix("www.")
def nco(s): return re.sub(r"[^a-z0-9]","",re.sub(r"\b(gmbh|inc|llc|ltd|limited|bv|b\.v\.|sas|sa|ag|ab|as|oy|aps|corp|co|plc|pty|the|kg)\b","",(s or "").lower()))
raw=[]
for f in sorted(glob.glob(os.path.join(D,"leads","*.jsonl"))):
    for i,l in enumerate(open(f)):
        l=l.strip()
        if not l: continue
        try: r=json.loads(l)
        except Exception: print("BADJSON",f,i); continue
        r["_f"]=os.path.basename(f); raw.append(r)
rank={"High":2,"Medium":1,"Low":0}
def score(r): return rank.get(r.get("research_confidence"),0)*2+(1 if r.get("email") else 0)
main=[];invalid=[];dups=0;idx={}
for r in raw:
    r={k:(v.strip() if isinstance(v,str) else v) for k,v in r.items()}
    r["country"]=CN.get(r.get("country",""),r.get("country",""))
    em=(r.get("email") or "").lower(); probs=[]
    if em and not EM.match(em): probs.append("invalid email format")
    if em and em.split("@")[1] in PERSONAL: probs.append("personal email domain")
    if not r.get("company"): probs.append("empty company")
    if not dom(r.get("website")): probs.append("no website")
    if r.get("country") not in OK_COUNTRIES: probs.append("country out of scope: %s"%r.get("country"))
    if not (r.get("source_url") or "").startswith("http"): probs.append("no source url")
    body=r.get("email_body","")
    if re.search(r"\[[^\]]+\]",body) or "<First" in body: probs.append("unresolved placeholder")
    if "Oseni Ibrahim" not in body: probs.append("missing signature")
    for fld in ("subject_line","email_body","personalized_opening"):
        for b in BANNED:
            if b in (r.get(fld) or "").lower(): probs.append("banned word '%s'"%b)
    if em in EXCL: probs.append(EXCL[em])
    if nco(r.get("company")) in EXCL: probs.append(EXCL[nco(r.get("company"))])
    if probs: r["_why"]="; ".join(probs); invalid.append(r); continue
    keys=[("c",nco(r["company"])),("w",dom(r["website"]))]
    if em: keys.append(("d",em.split("@")[1]))
    hit=next((idx[k] for k in keys if k[1] and k in idx),None)
    if hit is not None:
        dups+=1
        if score(r)>score(main[hit]): main[hit]=r
        for k in keys:
            if k[1]: idx[k]=hit
        continue
    for k in keys:
        if k[1]: idx[k]=len(main)
    main.append(r)
# status
for r in main:
    notes=[r.get("notes","")]
    ready=bool(r.get("email")) and r.get("research_confidence") in ("High","Medium") and r.get("personalization_detail")
    if not r.get("email"): notes.insert(0,"No verified email")
    lp=(r.get("email") or "").lower().split("@")[0]
    if lp in TECH:
        ready=False; notes.insert(0,"REVIEW: technical/registry inbox (%s@) - find a business contact before sending"%lp)
    for k,v in REVIEW.items():
        if k in (r.get("email") or "").lower() or k==nco(r.get("company")):
            ready=False; notes.insert(0,"REVIEW: "+v)
    r["_status"]="ready" if ready else "review"; r["_notes"]=" | ".join(n for n in notes if n)
print("raw",len(raw),"unique",len(main),"dups",dups,"invalid",len(invalid))
for r in invalid: print("INV",r["_f"],r.get("company"),"|",r["_why"])
sub=Counter(r["subject_line"] for r in main); print("dup subjects",[s for s,c in sub.items() if c>1][:10])
if "--write" not in sys.argv: sys.exit()
o={"High":0,"Medium":1,"Low":2}
main.sort(key=lambda r:(r["_status"]!="ready",o.get(r.get("buying_intent"),3),o.get(r.get("research_confidence"),3),r.get("country",""),r.get("company","").lower()))
wb=Workbook(); hf=PatternFill("solid",fgColor="1F3864"); thin=Side(style="thin",color="D9D9D9")
W=[24,24,12,18,30,14,14,14,24,18,34,18,14,16,12,28,28,10,40,45,45,45,34,70,40,55,20,11,10,10,10,40]
ws=wb.active; ws.title="Invoice Automation Prospects"; ws.append(COLS)
for r in main:
    ws.append([r.get(k,"") or "" for k in F]+[r["_status"],"Not Sent",None,r["_notes"]])
for i,w in enumerate(W,1):
    ws.column_dimensions[get_column_letter(i)].width=w
    c=ws.cell(1,i); c.font=Font(bold=True,color="FFFFFF"); c.fill=hf; c.alignment=Alignment(wrap_text=True,vertical="center",horizontal="center")
ws.row_dimensions[1].height=32
for rr in ws.iter_rows(min_row=2):
    for c in rr: c.alignment=Alignment(wrap_text=True,vertical="top"); c.border=Border(bottom=thin)
    ws.row_dimensions[rr[0].row].height=150
    rr[23].value=(rr[23].value or "").replace("\\n","\n")
    for j in (15,16,24):
        v=rr[j].value
        if v and str(v).startswith("http"): rr[j].hyperlink=v; rr[j].font=Font(color="0563C1",underline="single")
ws.freeze_panes="A2"
t=Table(displayName="InvoiceProspects",ref="A1:%s%d"%(get_column_letter(len(COLS)),max(ws.max_row,2)))
t.tableStyleInfo=TableStyleInfo(name="TableStyleLight9",showRowStripes=True); ws.add_table(t)
screened=int(sys.argv[sys.argv.index("--screened")+1]) if "--screened" in sys.argv else len(raw)
s=wb.create_sheet("Research Summary")
def block(title,counter):
    s.append((title,"")); s.cell(s.max_row,1).font=Font(bold=True)
    for k,v in counter.most_common(): s.append(("  "+(k or "(blank)"),v))
    s.append(("",""))
s.append(("Metric","Value"))
for c in s[1]: c.font=Font(bold=True,color="FFFFFF"); c.fill=hf
for k,v in [("Total prospects (unique)",len(main)),("Candidates screened (approx.)",screened),("Raw research records",len(raw)),
    ("Prospects with verified emails",sum(1 for r in main if r.get("email"))),("Prospects without verified emails",sum(1 for r in main if not r.get("email"))),
    ("Lead Status ready",sum(r["_status"]=="ready" for r in main)),("Lead Status review",sum(r["_status"]=="review" for r in main)),
    ("Duplicate count removed",dups),("Invalid prospects removed",len(invalid))]: s.append((k,v))
s.append(("",""))
block("Prospects by country",Counter(r["country"] for r in main))
block("Prospects by company type",Counter(r.get("company_type") for r in main))
block("Prospects by buying intent",Counter(r.get("buying_intent") for r in main))
block("Prospects by research confidence",Counter(r.get("research_confidence") for r in main))
s.append(("Verification method","Company websites, LinkedIn and Google were blocked by the research environment's network policy. Facts were verified on GitHub organization/repository pages, npm and PyPI package metadata, and search results quoting company pages. Emails come only from those public sources; none were guessed."))
s.append(("Send status","No emails sent. Send Status = Not Sent, Date Sent blank for all rows."))
s.column_dimensions["A"].width=48; s.column_dimensions["B"].width=100
for rr in s.iter_rows(min_row=2):
    for c in rr: c.alignment=Alignment(wrap_text=True,vertical="top")
wb.save(OUT); print("saved",OUT)
