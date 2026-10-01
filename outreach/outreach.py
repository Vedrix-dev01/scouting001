"""Outreach pacing + tracking.  Usage:
  python3 outreach/outreach.py next                      -> JSON of next email to send, or {"wait_until": ISO}
  python3 outreach/outreach.py record <idx> <msgId> <threadId>
  python3 outreach/outreach.py mark <email> <replied|optout|bounced> [note]
  python3 outreach/outreach.py status
"""
import json, sys, os, random
from datetime import datetime, timedelta, timezone
from zoneinfo import ZoneInfo
D=os.path.dirname(os.path.abspath(__file__)); R=os.path.dirname(D)
Q=json.load(open(os.path.join(D,"queue.json")))
SP=os.path.join(D,"state.json")
S=json.load(open(SP)) if os.path.exists(SP) else {"start":"2026-10-01","sent":{},"status":{}}
TZ={"United States":"America/Chicago","Canada":"America/Toronto","Australia":"Australia/Sydney","United Kingdom":"Europe/London",
    "Ireland":"Europe/Dublin","Portugal":"Europe/Lisbon","Iceland":"Atlantic/Reykjavik","Finland":"Europe/Helsinki","Estonia":"Europe/Tallinn",
    "Lithuania":"Europe/Vilnius","Romania":"Europe/Bucharest","Cyprus":"Asia/Nicosia","Germany":"Europe/Berlin"}
FOOT={"en":"\n\nIf you'd prefer not to hear from me, just reply \"no thanks\" and I won't email again.",
      "de":"\n\nFalls Sie keine weiteren Nachrichten wünschen, antworten Sie einfach kurz mit „Nein danke“."}
def save(): json.dump(S,open(SP,"w"),indent=1)
def now(): return datetime.now(timezone.utc)
def cap(t):
    d=(t.date()-datetime.fromisoformat(S["start"]).date()).days
    return 20 if d<7 else 30 if d<14 else 40
def sent_today(t): return sum(1 for v in S["sent"].values() if v["sent_at"][:10]==t.strftime("%Y-%m-%d"))
def open_(item,t):
    lt=t.astimezone(ZoneInfo(TZ.get(item["country"],"Europe/Berlin")))
    return lt.weekday()<5 and 9<=lt.hour<17
def pending(): return [(i,q) for i,q in enumerate(Q) if str(i) not in S["sent"] and S["status"].get(q["email"].lower()) not in ("optout","bounced")]
cmd=sys.argv[1]
if cmd=="next":
    t=now(); p=pending()
    if not p: print(json.dumps({"done":True})); sys.exit()
    if sent_today(t)>=cap(t):
        nxt=(t+timedelta(days=1)).replace(hour=0,minute=5,second=0,microsecond=0)
        print(json.dumps({"wait_until":nxt.isoformat(),"reason":"daily cap reached"})); sys.exit()
    for i,q in p[:400]:
        if open_(q,t):
            print(json.dumps({"idx":i,"to":q["email"],"company":q["company"],"country":q["country"],"subject":q["subject"],"body":q["body"]+FOOT[q["lang"]],"delay_after":random.randint(20,35)},ensure_ascii=False)); sys.exit()
    for k in range(1,4*24*4):
        tt=t+timedelta(minutes=15*k)
        if any(open_(q,tt) for _,q in p[:400]):
            print(json.dumps({"wait_until":tt.isoformat(),"reason":"outside business hours"})); sys.exit()
elif cmd=="record":
    i,mid,tid=int(sys.argv[2]),sys.argv[3],sys.argv[4]; q=Q[i]; t=now()
    S["sent"][str(i)]={"email":q["email"],"company":q["company"],"msg":mid,"thread":tid,"sent_at":t.isoformat()}; save()
    from openpyxl import load_workbook
    wb=load_workbook(os.path.join(R,q["file"])); ws=wb[q["sheet"]]; H=[c.value for c in ws[1]]
    ws.cell(q["row"],H.index("Send Status")+1).value="Sent"; ws.cell(q["row"],H.index("Date Sent")+1).value=t.strftime("%Y-%m-%d %H:%M UTC")
    wb.save(os.path.join(R,q["file"])); print("recorded",q["company"])
elif cmd=="mark":
    em,st=sys.argv[2].lower(),sys.argv[3]; S["status"][em]=st; save()
    from openpyxl import load_workbook
    for q in Q:
        if q["email"].lower()==em:
            wb=load_workbook(os.path.join(R,q["file"])); ws=wb[q["sheet"]]; H=[c.value for c in ws[1]]
            ws.cell(q["row"],H.index("Send Status")+1).value={"replied":"Replied","optout":"Opted out","bounced":"Bounced"}.get(st,st)
            n=ws.cell(q["row"],H.index("Notes")+1); n.value=((n.value or "")+f" | {st} {now():%Y-%m-%d}"+(" - "+sys.argv[4] if len(sys.argv)>4 else "")).strip(" |")
            wb.save(os.path.join(R,q["file"]))
    print("marked",em,st)
elif cmd=="status":
    t=now(); print(json.dumps({"queued":len(Q),"sent":len(S["sent"]),"sent_today":sent_today(t),"cap_today":cap(t),"statuses":S["status"]},indent=1))
